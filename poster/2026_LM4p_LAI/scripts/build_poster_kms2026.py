#!/usr/bin/env python3
"""KMS 2026 (한국기상학회 가을학술대회) poster, Korean text, 90 x 140 cm, same template as the ACM poster.

Focus (user decision 2026-10-08): global LAI evaluation (tropical / subtropical excess, SE South America) +
model species distribution vs an observation-based map; spin-up figure in Methods; phenology only one line
(the pilot experiments are on the ACM poster). Helpers copied from build_poster_acm2026.py.
template : /Volumes/data01/발표_템플릿/전남대/poster/poster_2026_ACM.pptx (never modified)
output   : ../poster_2026_KMS_draft.pptx
Numbers: data/kms_zonal_month_bias.txt (p16), data/kms_species_shares.txt (p17), data/spinup_cycle_drift.txt (p0),
         lm4/LM4_SPINUP_NOTES.md 13.66/13.70 (species census, MJJAS +41 %).
"""
import copy
import shutil
from pptx import Presentation
from pptx.util import Pt

T = "/Volumes/data01/발표_템플릿/전남대/poster/poster_2026_ACM.pptx"
OUT = "../poster_2026_KMS_draft.pptx"
FIG = "../figures/"
shutil.copy2(T, OUT)
p = Presentation(OUT); s = p.slides[0]
NAMES = {"Header: RESULTS 1": "Header: RESULTS 1: OBSERVED RELATIONSHIP", "Header text: RESULTS 1": "Header text: RESULTS 1: OBSERVED RELATIONSHIP",
         "Header: RESULTS 2": "Header: RESULTS 2: MECHANISM", "Header text: RESULTS 2": "Header text: RESULTS 2: MECHANISM",
         "Header: RESULTS 3": "Header: RESULTS 3: MODEL EXPERIMENT", "Header text: RESULTS 3": "Header text: RESULTS 3: MODEL EXPERIMENT",
         "Header: SUMMARY": "Header: SUMMARY & CONCLUSIONS", "Header text: SUMMARY": "Header text: SUMMARY &amp; CONCLUSIONS",
         "Header: REFERENCES": "Header: REFERENCES & ACKNOWLEDGEMENTS", "Header text: REFERENCES": "Header text: REFERENCES &amp; ACKNOWLEDGEMENTS",
         "Header: INTRO": "Header: INTRODUCTION", "Header text: INTRO": "Header text: INTRODUCTION",
         "Header: DATA": "Header: DATA & METHODS", "Header text: DATA": "Header text: DATA &amp; METHODS",
         "Figure 1 placeholder": "Figure 1 placeholder box", "Figure 2 placeholder": "Figure 2 placeholder box", "Figure 3 placeholder": "Figure 3 placeholder box"}
_by_name = {x.name: x for x in s.shapes if x.name != "Text 1"}
sh = {k: _by_name[NAMES.get(k, k)] for k in list(NAMES) + ["Introduction body", "Data table", "Methods body", "Figure 1 placeholder label", "Caption 1",
      "Results 1 body", "Figure 2 placeholder label", "Caption 2", "Results 2 body", "Figure 3 placeholder label", "Results 3 body", "Caption 3",
      "Conclusions body", "References body", "TextBox 52"]}
titles = [x for x in s.shapes if x.name == "Text 1"]            # title, authors, affiliations (template order)
IN = 914400
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
CHECK_PPR = ('<a:pPr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" marL="457200" indent="-457200">'
             '<a:buFont typeface="Wingdings" pitchFamily="2" charset="2"/><a:buChar char="&#252;"/></a:pPr>')


def fill(box, items, size=None, bullet=None):
    from lxml import etree
    tf = box.text_frame
    p0 = tf.paragraphs[0]._p; pPr = p0.find(A + "pPr")
    r0 = tf.paragraphs[0].runs[0]._r; rPr = r0.find(A + "rPr")
    if bullet:
        pPr = etree.fromstring(bullet)
        old = p0.find(A + "pPr")
        if old is not None:
            p0.remove(old)
        p0.insert(0, copy.deepcopy(pPr))
    for par in list(tf.paragraphs)[1:]:
        par._p.getparent().remove(par._p)
    for k, (text, bold) in enumerate(items):
        par = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        if k > 0 and pPr is not None:
            par._p.insert(0, copy.deepcopy(pPr))
        for r in list(par.runs):
            r._r.getparent().remove(r._r)
        run = par.add_run(); run.text = text
        if rPr is not None:
            run._r.insert(0, copy.deepcopy(rPr))
        run.font.bold = bold
        if size:
            run.font.size = Pt(size)


def header(key, text, top, size=None):
    for nm in ("Header: " + key, "Header text: " + key):
        sh[nm].top = int(top * IN)
    r = sh["Header text: " + key].text_frame.paragraphs[0].runs[0]; r.text = text
    if size:
        r.font.size = Pt(size)


def picture(png, top, width, left_col=True, box=None, label=None):
    x0 = 1.0 if left_col else 18.1
    pic = s.shapes.add_picture(FIG + png, int((x0 + (16.3 - width) / 2) * IN), int(top * IN), width=int(width * IN))
    for b in (box, label):
        if b is not None:
            b._element.getparent().remove(b._element)
    pic.name = "Picture " + png
    return pic


def put(box, top, height, left_col=True, width=16.3):
    box.left = int((1.0 if left_col else 18.1) * IN); box.width = int(width * IN)
    box.top = int(top * IN); box.height = int(height * IN)
    from pptx.enum.text import MSO_ANCHOR
    box.text_frame.vertical_anchor = MSO_ANCHOR.TOP      # template boxes are middle-anchored; top keeps text under its caption


# ---------------------------------------------------------------- title block
sh["TextBox 52"].text_frame.paragraphs[0].runs[0].text = "2026 한국기상학회 가을학술대회"
titles[0].text_frame.paragraphs[0].runs[0].text = "LM4+ 지면 모델의 전 지구 엽면적지수 모의 특성 진단"
ar = titles[1].text_frame.paragraphs[0].runs
for r, t in zip(ar, ["고영대", "¹", ", ", "박상혁", "", "², ", "정지훈", "²"]):
    r.text = t
af = titles[2].text_frame.paragraphs
af[0].runs[0].text = "¹전남대학교 해양학과"
af[1].runs[0].text = "²세종대학교 환경융합공학과"
af[2].runs[0].text = "Contact: youngdaekoh@gmail.com"

# ---------------------------------------------------------------- left column
y = 11.3
header("INTRO", "서론", y); y += 1.5
put(sh["Introduction body"], y, 6.6)
fill(sh["Introduction body"], [
    ("한국해양과학기술원(KIOST)은 CMIP7 참여를 위해 KIOST-ESM2를 개발 중이며, 지면 성분으로 GFDL의 LM4+를 사용한다.", True),
    ("LM4+는 식생을 크기별 개체군(cohort, PPA)으로 모의하여 엽면적·생물량·종 구성이 기후에 반응해 변한다. "
     "이 자유도는 처방 식생이 담지 못하는 되먹임을 표현하지만, 동시에 큰 불확실성의 원인이 될 수 있다.", False),
    ("따라서 완전 결합 실험에 앞서, 관측 강제력 하의 오프라인 실험으로 지면 모델 자체의 식생 편차를 진단할 필요가 있다.", False),
    ("목표: 1982–2010년 LM4+ 엽면적지수(LAI)를 위성 관측(GIMMS LAI4g)과 비교하여 편차의 위치·시기·원인(종 분포 포함)을 진단한다.", True),
], size=28)
y += 6.8
header("DATA", "자료 및 방법", y); y += 1.5
tbl = sh["Data table"].table
tbl._tbl.append(copy.deepcopy(tbl._tbl.tr_lst[-1]))
rows = [["자료", "설정 / 변수", "기간 / 해상도"],
        ["LM4+ 오프라인 (KIOST-ESM2 지면)", "PPA cohort, 6개 종, 종 분포 고정, 질소 순환 off, CO₂ 고정", "1979–2023, C96 (≈1°) → 1°"],
        ["WFDE5 강제력 (Cucchi et al., 2020)", "ERA5를 CRU/GPCC로 편향 보정: 기온, 습도, 바람, 단파, 장파, 강수", "0.5°, 1시간 → C96"],
        ["초기장", "KIOST-ESM2 약 400년 결합 piControl + WFDE5 1981–2010 1회 spin-up", "1979-01-01 시작"],
        ["GIMMS LAI4g (Cao et al., 2023)", "AVHRR LAI, MODIS로 통합", "1982–2010, 15일 1/12° → 월 1°"]]
for i, row in enumerate(rows):
    tbl.rows[i].height = int((0.75 if i == 0 else 1.0) * IN)
    for j, txt in enumerate(row):
        par = tbl.cell(i, j).text_frame.paragraphs[0]
        r0 = par.runs[0] if par.runs else par.add_run()
        for r in list(par.runs)[1:]:
            r._r.getparent().remove(r._r)
        r0.text = txt; r0.font.size = Pt(21)
sh["Data table"].top = int(y * IN)
y += 4.9
pic0 = picture("fig0_spinup_lai_gpp_wood_kms.png", y, 15.8, box=sh["Figure 1 placeholder"], label=sh["Figure 1 placeholder label"])
y += pic0.height / IN + 0.1
put(sh["Methods body"], y, 2.6)
fill(sh["Methods body"], [
    ("Fig. 1. Offline spin-up after the KIOST-ESM2 initial state (WFDE5 1981–2010 twice; non-ice land, annual mean): "
     "(a) LAI, (b) GPP, (c) wood carbon. Shading: cycle 2 = analysis run.", True),
    ("두 번째 주기 변화: LAI −0.6 %, GPP +0.4 % → 평형, 느린 목재 탄소만 +7.6 % 계속 증가. "
     "평가: 1982–2010 월 기후값을 1° 공통 격자에서 비교 (MJJAS 평균, 위도·월별 편차, 종 분포).", False),
], size=22)
y += 2.8
header("RESULTS 1", "결과 1: 전 지구 LAI 분포", y); y += 1.5
pic1 = picture("fig1_lai_mean_bias.png", y, 11.5)
y += pic1.height / IN + 0.1
put(sh["Caption 1"], y, 1.3)
fill(sh["Caption 1"], [("Fig. 2. Growing-season (MJJAS) LAI, 1982–2010: (a) GIMMS LAI4g, (b) LM4+, (c) LM4+ − GIMMS. Side panels: zonal means.", True)], size=24)
y += 1.4
put(sh["Results 1 body"], y, 55.0 - y)
fill(sh["Results 1 body"], [
    ("공간 분포는 재현되나 거의 전 지역에서 과대: 전 지구 MJJAS 평균 2.77 vs 1.97 (+41 %).", True),
    ("남아메리카 남동부(AR6 SES)에서 가장 뚜렷: 1.41 → 5.17 (+268 %). 열대 남미·아프리카·동남아도 +30–49 %.", True),
    ("한대(50–70°N)는 여름 정점 +15 %뿐, 편차의 87 %가 비생장기(9–4월)에서 발생 → 크기보다 계절 위상 문제.", False),
], size=26)

# ---------------------------------------------------------------- right column
y = 11.3
header("RESULTS 2", "결과 2: 편차는 언제, 어디서?", y); y += 1.5
pic2 = picture("fig_kms_zonal_month_bias.png", y, 15.6, left_col=False, box=sh["Figure 2 placeholder"], label=sh["Figure 2 placeholder label"])
y += pic2.height / IN + 0.1
put(sh["Caption 2"], y, 1.3, left_col=False)
fill(sh["Caption 2"], [("Fig. 3. Latitude–month zonal means, 1982–2010: (a) GIMMS LAI, (b) LM4+ − GIMMS. Cells with GIMMS annual-mean LAI < 0.25 (deserts, polar barrens) are excluded.", True)], size=24)
y += 1.4
put(sh["Results 2 body"], y, 4.3, left_col=False)
fill(sh["Results 2 body"], [
    ("열대(20°S–20°N): 연중 +1.2 (+42 %), 계절과 무관한 크기 편차.", False),
    ("남반구 중위도(20–60°S): 연중 +1.5 (+112 %). 관측 LAI가 낮은 남미 남동부(1.4)에서 모델은 숲 수준(5.2).", True),
    ("한대(50–70°N): 여름 +0.30이지만 가을(SON) +1.12 → 낙엽 지연. 원인은 기온 단일 문턱(평활 기온 < 5 °C)의 낙엽 규칙 (Weng et al., 2015).", False),
], size=26)
y += 4.5
header("RESULTS 3", "결과 3: 관측 기반 식생 지도 vs 모델 종 분포", y); y += 1.5
pic3 = picture("fig_kms_species_model_vs_obs.png", y, 16.0, left_col=False, box=sh["Figure 3 placeholder"], label=sh["Figure 3 placeholder label"])
y += pic3.height / IN + 0.1
put(sh["Caption 3"], y, 1.3, left_col=False)
fill(sh["Caption 3"], [("Fig. 4. Dominant LM4+ species per 1° cell: (a) observation-based (MODIS-derived CLM5 PFT map, "
                        "Lawrence et al., 2019), (b) LM4+ initial state from the KIOST-ESM2 spin-up (July leaf area).", True)], size=24)
y += 1.4
put(sh["Results 3 body"], y, 4.6, left_col=False)
fill(sh["Results 3 body"], [
    ("열대(30°S–30°N): C4 풀 관측 45 % vs 모델 8 %, C3 풀 17 % vs 48 % → 사바나의 C4 풀이 C3 풀·열대 상록수로 대체.", True),
    ("50°N 이북: 낙엽송 7.5 % vs 0 %, 낙엽 활엽수 4.8 % vs 16.7 %. 낙엽송은 결합 spin-up 경쟁에서 북반구에서 사라짐 (MODIS 원자료에서도 동시베리아 낙엽 침엽수 확인).", True),
    ("종 분포가 고정된 설정에서는 이 편차가 그대로 유지 → 관측 지도로 다시 심은 초기장을 시험 중.", False),
], size=26)
y += 4.8
header("SUMMARY", "요약 및 향후 계획", y); y += 1.5
put(sh["Conclusions body"], y, 5.0, left_col=False)
fill(sh["Conclusions body"], [
    ("LM4+는 전 지구 LAI 분포를 재현하나 MJJAS +41 % 과대, 특히 남미 남동부(+268 %)와 남반구 중위도 초지.", True),
    ("편차는 열대·남반구의 연중 크기 과대와 한대의 가을 낙엽 지연(계절 위상)으로 구분된다.", True),
    ("모델 종 분포는 관측과 달리 열대 C4 풀이 거의 없고, 북방 낙엽송이 사라졌으며 낙엽 활엽수가 과다하다.", True),
    ("향후: 낙엽 규칙(낮 길이 조건) 개선, 관측 지도 기반 재파종, 시간 변화 CO₂, KIOST-ESM2 결합 실험에서의 영향 평가.", False),
], size=24, bullet=CHECK_PPR)
y += 5.2
header("REFERENCES", "참고문헌 및 사사", y); y += 1.4
put(sh["References body"], y, 55.0 - y, left_col=False)
fill(sh["References body"], [
    ("Cao, S., et al. (2023). ESSD, 15, 4877. • Cucchi, M., et al. (2020). ESSD, 12, 2097. • Shevliakova, E., et al. (2024). JAMES, 16. "
     "• Weng, E., et al. (2015). Biogeosciences, 12, 2655. • Iturbide, M., et al. (2020). ESSD, 12, 2959. "
     "• Lawrence, D. M., et al. (2019). JAMES, 11, 4245. • Lawrence, P. J. & Chase, T. N. (2007). JGR, 112, G01023.", False),
    ("이 연구는 기후에너지환경부의 재원으로 한국환경산업기술원의 “신기후체제 대응 환경기술개발사업”의 지원(RS-2022-KE002160) 및 "
     "해양수산부의 재원으로 해양수산과학기술진흥원의 지원(RS-2026-25543619)을 받아 수행되었습니다.", True),
], size=20)
p.save(OUT)
print("saved", OUT, "| left column ends at %.1f in, right column ends at %.1f in (slide 55.1)" % (
    sh["Results 1 body"].top / IN + sh["Results 1 body"].height / IN, y))
for nm in [x.name for x in s.shapes if x.name.startswith("Picture")]:
    x = [q for q in s.shapes if q.name == nm][0]
    print("  %-46s y %.1f..%.1f  w %.1f" % (nm, x.top / IN, (x.top + x.height) / IN, x.width / IN))
