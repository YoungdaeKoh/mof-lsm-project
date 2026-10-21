# -*- coding: utf-8 -*-
"""Build 1차년도 착수보고회 deck from template (세부 1.7, 세종대 정지훈)."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import MSO_CONNECTOR

SRC = "차세대_착수보고회_템플릿.pptx"
OUT = "차세대_착수보고회_세종대_1.7_1차년도.pptx"

NAVY = "0F304B"
NAVY2 = "10314B"
ACCENT = "1F6FB2"
WHITE = "FFFFFF"
GREY = "5A6B7A"

prs = Presentation(SRC)

# ---------- helpers ----------
def dup_slide(prs, src):
    dest = prs.slides.add_slide(src.slide_layout)
    for sh in list(dest.shapes):
        sh._element.getparent().remove(sh._element)
    rId_map = {}
    for rId, rel in src.part.rels.items():
        if rel.reltype.endswith("/slideLayout"):
            continue
        if rel.is_external:
            new = dest.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            new = dest.part.relate_to(rel.target_part, rel.reltype)
        rId_map[rId] = new
    embed = qn('r:embed'); link = qn('r:id'); rlink = qn('r:link')
    for sh in src.shapes:
        el = copy.deepcopy(sh._element)
        for node in el.iter():
            for attr in (embed, rlink):
                v = node.get(attr)
                if v in rId_map:
                    node.set(attr, rId_map[v])
        dest.shapes._spTree.append(el)
    return dest

def find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None

def set_keep(shape, text):
    """Replace text keeping first run formatting."""
    tf = shape.text_frame
    p = tf.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for rr in p.runs[1:]:
            rr._r.getparent().remove(rr._r)
    else:
        p.text = text
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)

def set_title(slide, text, size=38):
    tb = find(slide, "TextBox 6")
    if tb:
        set_keep(tb, text)
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size)

def del_placeholder(slide):
    tb = find(slide, "TextBox 2")
    if tb is not None:
        tb._element.getparent().remove(tb._element)

def set_pagenum(slide, s):
    tb = find(slide, "TextBox 7")
    if tb is not None:
        set_keep(tb, s)

def add_body(slide, lines, left=1.55, top=2.55, width=23.6, height=11.7):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for spec in lines:
        txt, sz, bold, color, sb = spec
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        run = p.add_run(); run.text = txt
        f = run.font; f.name = "Pretendard"; f.size = Pt(sz); f.bold = bold
        f.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(sb)
        p.line_spacing = 1.06
    return tb

def H(t, sz=30):   return (t, sz, True, NAVY, 5)      # header
def I(t, sz=25):   return ("•  " + t, sz, False, NAVY, 4)   # item
def S(t, sz=21):   return ("      –  " + t, sz, False, GREY, 3)  # sub
def B(sz=8):       return ("", sz, False, NAVY, 0)    # blank spacer

def style_table(tbl, header_sz=20, body_sz=19, col0_bold=True):
    for ri, row in enumerate(tbl.rows):
        for ci, cell in enumerate(row.cells):
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.12)
            cell.margin_top = Inches(0.05); cell.margin_bottom = Inches(0.05)
            for p in cell.text_frame.paragraphs:
                p.line_spacing = 1.0
                for r in p.runs:
                    r.font.name = "Pretendard"
                    if ri == 0:
                        r.font.size = Pt(header_sz); r.font.bold = True
                        r.font.color.rgb = RGBColor.from_string(WHITE)
                    else:
                        r.font.size = Pt(body_sz)
                        r.font.bold = bool(col0_bold and ci == 0)
                        r.font.color.rgb = RGBColor.from_string(NAVY)
            if ri == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(NAVY)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string("EAF0F5" if ri % 2 else "FFFFFF")

def rrect(slide, x, y, w, h, fill, text="", sz=14, bold=False,
          fcolor=NAVY, line=None, align=PP_ALIGN.CENTER, radius=0.08):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sp.adjustments[0] = radius
    sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sp.line.color.rgb = RGBColor.from_string(line); sp.line.width = Pt(1)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.03); tf.margin_bottom = Inches(0.03)
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = 1.02
        r = p.add_run(); r.text = ln
        r.font.name = "Pretendard"; r.font.size = Pt(sz); r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(fcolor)
    return sp

def vline(slide, x, y0, y1, color="C9D3DD"):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    Inches(x), Inches(y0), Inches(x), Inches(y1))
    cn.line.color.rgb = RGBColor.from_string(color); cn.line.width = Pt(1.0)
    cn.shadow.inherit = False
    return cn

def add_table(slide, data, left, top, width, height, col_w=None):
    rows, cols = len(data), len(data[0])
    gframe = slide.shapes.add_table(rows, cols, Inches(left), Inches(top),
                                    Inches(width), Inches(height))
    tbl = gframe.table
    for r in range(rows):
        for c in range(cols):
            tbl.cell(r, c).text = data[r][c]
    if col_w:
        for c, w in enumerate(col_w):
            tbl.columns[c].width = Inches(w)
    return tbl

# ---------- references ----------
s_title = prs.slides[0]
s_toc   = prs.slides[1]
s_plan  = prs.slides[2]
s_y1    = prs.slides[3]
s_perf  = prs.slides[4]
s_sugg  = prs.slides[5]
s_thx   = prs.slides[6]

plan_b = dup_slide(prs, s_plan)
plan_c = dup_slide(prs, s_plan)
y1_b   = dup_slide(prs, s_y1)
y1_c   = dup_slide(prs, s_y1)
y1_d   = dup_slide(prs, s_y1)

# ================= TITLE =================
set_keep(find(s_title, "TextBox 20"), "지면-식생모델 성능 및 해양영향 진단 체계 구축 (세부 1.7)")
set_keep(find(s_title, "TextBox 18"), "세종대학교 산학협력단")
set_keep(find(s_title, "TextBox 19"), "정지훈")

# ================= 전체계획 A : 연구 배경/필요성 =================
del_placeholder(s_plan); set_title(s_plan, "연구 배경 및 필요성")
add_body(s_plan, [
    H("■ 연구 배경"),
    I("지면-식생모델은 육상 물·에너지 순환 모의의 핵심 → 기후모델 정확도 좌우"),
    I("지면-식생 변화가 지면-대기 상호작용으로 열·수분·에너지 수지 조절,"),
    S("대기순환 원격상관을 통해 해양·전지구 기후시스템에 영향"),
    I("급격한 기후변화에 따른 지면·유출량 변동 → 연안 해양의 염분·성층화에 직접 영향"),
    B(),
    H("■ 연구 필요성"),
    I("구조적 오차: 열·수분 플럭스 모의 오차가 수문·유출 불확실성을 가중 → 해양모델로 전이"),
    I("다중 모델 벤치마킹 시급: 결합모델에 탑재된 지면모델 한계의 객관적 진단 필요"),
    I("물리과정 고도화: 진단 기반 수문·역학식생·담수유출 모델 개선 → 한반도 주변 기후모의 최적화"),
])

# ================= 전체계획 B : 목표/추진전략 =================
del_placeholder(plan_b); set_title(plan_b, "연구 목표 및 추진 전략")
add_body(plan_b, [
    H("■ 최종 목표 (2026–2030)"),
    I("차세대 해양기후모델용 지면-식생모델 진단·개선 체계 구축"),
    I("세부 1.7 역할: 지면→대기→해양 상호작용 진단 (담수 유출이 지면-해양 결합 매개)"),
    B(),
    H("■ 추진 전략 (4단계)"),
    I("① 다중 LSM 오프라인 테스트베드 구축"),
    I("② 관측·재분석 기반 기본성능 진단"),
    I("③ 구조오차·기후변화 민감도 분석"),
    I("④ 물리과정 개선 및 지면-해양 결합 고도화"),
    B(),
    H("■ 단계 구성"),
    I("Phase 1 (2026–2028): 진단 체계 구축 및 개선점 도출"),
    I("Phase 2 (2029–2030): 물리과정 고도화 및 미래 해양기후 영향 예측"),
])

# ================= 전체계획 C : 5개년 로드맵 (연차별 매트릭스) =================
del_placeholder(plan_c); set_title(plan_c, "연차별 상세 연구 내용 (5개년 로드맵)")
# colors
BLUE = "2F5597"; PEACH = "F4B183"; PEACH_LN = "C55A11"
BAR_DATA = "ED7D31"; BAR_DIAG = "FFC000"; BAR_MODEL = "70AD47"
LGREY = "EDF1F5"
# geometry
GX0 = 3.30; GW = 22.90               # column area
NCOL = 5; GAP = 0.18
COLW = (GW - (NCOL - 1) * GAP) / NCOL
def cx(c): return GX0 + c * (COLW + GAP)
LBL_X = 0.40; LBL_W = 2.65; BAR_X = 0.40; BAR_W = 0.20
# top goal box
rrect(plan_c, GX0, 2.45, GW, 1.15,
      "10314B",
      "지면-식생-대기-해양 상호작용을 진단·개선하여\n"
      "차세대 해양기후모델의 한반도 주변 기후모의 성능 향상",
      sz=22, bold=True, fcolor=WHITE, radius=0.06)
# year labels + header boxes
YEARS = ["2026", "2027", "2028", "2029", "2030"]
HEAD = ["진단 체계 구축", "기후변화 민감도\n· 구조오차", "CMIP7 상호작용\n진단",
        "물리과정·유출\n고도화", "전망·예측성능\n평가"]
HEAD_Y = 4.45; HEAD_H = 1.15
for c in range(NCOL):
    tb = plan_c.shapes.add_textbox(Inches(cx(c)), Inches(4.02), Inches(COLW), Inches(0.42))
    pp = tb.text_frame.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    rr = pp.add_run(); rr.text = YEARS[c]
    rr.font.name = "Pretendard"; rr.font.size = Pt(18); rr.font.bold = True
    rr.font.italic = True; rr.font.color.rgb = RGBColor.from_string(NAVY)
    rrect(plan_c, cx(c), HEAD_Y, COLW, HEAD_H, BLUE, HEAD[c],
          sz=16, bold=True, fcolor=WHITE, radius=0.10)
# row bands: (label, bar color, y, h)
ROWS = [
    ("자료구축\n· 테스트베드", BAR_DATA, 5.95, 2.30),
    ("진단\n· 민감도", BAR_DIAG, 8.40, 2.70),
    ("모델 개선\n· 고도화", BAR_MODEL, 11.20, 1.75),
]
for label, barc, ry, rh in ROWS:
    rrect(plan_c, BAR_X, ry, BAR_W, rh, barc, radius=0.20)
    tb = plan_c.shapes.add_textbox(Inches(0.72), Inches(ry), Inches(LBL_W), Inches(rh))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(label.split("\n")):
        pp = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pp.alignment = PP_ALIGN.LEFT
        rr = pp.add_run(); rr.text = ln
        rr.font.name = "Pretendard"; rr.font.size = Pt(18); rr.font.bold = True
        rr.font.color.rgb = RGBColor.from_string(NAVY)
# vertical guide lines between columns
for c in range(NCOL + 1):
    xx = cx(c) - GAP / 2 if c < NCOL else cx(NCOL - 1) + COLW + GAP / 2
    vline(plan_c, xx, HEAD_Y, 12.95)
# cells: keyed by (row_idx, col_idx) -> list of box texts
CELLS = {
    (0, 0): ["4종 LSM(JULES·CLM5·Noah-MP·LM4+)\n오프라인 테스트베드 구축",
             "관측·재분석·기후프록시 자료 DB 구축"],
    (0, 2): ["CMIP7 시나리오\n자료 확보·전처리"],
    (1, 0): ["1979–현재 offline 적분,\n토양수분/온도·식생 기본성능 진단"],
    (1, 1): ["산업화이후 장기 민감도 실험",
             "결합강도·수문·유출 민감도 /\n관측대비 구조오차 진단"],
    (1, 2): ["CMIP7 지면-대기 결합강도·\n해양영향 진단, 개선점 도출"],
    (1, 4): ["극한기후(폭염-가뭄)·\n북서태평양 해양열파\n예측성능 평가"],
    (2, 3): ["수문 물리과정·담수\n유출모델 업데이트",
             "물리과정 AI 에뮬레이터\n적용 시도"],
    (2, 4): ["역학식생 모델 업데이트",
             "미래 담수유출→동아시아\n연안 영향 전망"],
}
PAD = 0.10
for (ri, ci), boxes in CELLS.items():
    _, _, ry, rh = ROWS[ri]
    x = cx(ci) + 0.04; w = COLW - 0.08
    n = len(boxes); inner = rh - 2 * PAD; gap = 0.12
    bh = (inner - (n - 1) * gap) / n
    for k, txt in enumerate(boxes):
        by = ry + PAD + k * (bh + gap)
        rrect(plan_c, x, by, w, bh, PEACH, txt, sz=13.5, bold=False,
              fcolor="3A2A12", line=PEACH_LN, radius=0.10)

# ================= 1차년도 A : 추진 개요 =================
del_placeholder(s_y1); set_title(s_y1, "1차년도(2026) 추진 개요")
add_body(s_y1, [
    H("■ 1차년도 목표: 지면-식생모델 성능·해양영향 진단 체계 구축"),
    B(),
    H("■ 3대 추진 과제 (점검항목)"),
    I("① 테스트베드  —  오프라인 지면-식생모델(벤치마크) 구축"),
    I("② 자료구축  —  평가용 토양수분/온도·식생 관측·재분석 DB 구축"),
    I("③ 성능진단  —  1979–현재 offline simulation, 모델-관측 비교"),
    B(),
    H("■ 추진 흐름"),
    I("ERA5 forcing → ① 다중 LSM 장기적분 → ② 관측 DB 대비 검증 →"),
    S("③ 토양수분/온도·식생 기본성능 진단 + 지면-대기·해양영향 기초 진단도구"),
])

# ================= 1차년도 B : ① 테스트베드 =================
del_placeholder(y1_b); set_title(y1_b, "① 오프라인 지면-식생모델 테스트베드")
add_body(y1_b, [
    H("■ 다중 LSM 벤치마크 체계 (4종)"),
    I("JULES vn7.4 · CESM2-CLM5 · Noah-MP · LM4+"),
    I("공통 forcing·격자·진단 프로토콜로 모델 간 객관적 비교"),
    B(),
    H("■ 구축 현황 (2026.06 기준)"),
    I("JULES vn7.4: 빌드·테스트런 완료, spin-up 설정 진행"),
    I("CESM2-CLM5: 빌드 완료, offline(I2000Clm50Sp) 설정"),
    I("Noah-MP · LM4+: 구축 예정"),
    B(),
    H("■ 실행 환경"),
    I("climate00 HPC (48 cores, Intel oneAPI, NetCDF/HDF5/PnetCDF)"),
])

# ================= 1차년도 C : ② 자료 DB =================
del_placeholder(y1_c); set_title(y1_c, "② 평가용 관측·재분석 자료 DB")
add_body(y1_c, [
    H("■ 모델 입력 (Forcing)"),
    I("ERA5 재분석 (1979–현재) → 모델별 forcing 변환기 구축 (JULES/CLM5/Noah-MP)"),
    B(),
    H("■ 검증용 관측·재분석 변수"),
    I("토양수분: GLEAM · ESA CCI · ERA5-Land"),
    I("토양온도: ERA5-Land · FLUXNET"),
    I("식생(LAI/GPP): MODIS · FLUXNET"),
    I("유출/물저장: GRDC · GRACE"),
    B(),
    H("■ 벤치마킹 도구"),
    I("ILAMB · PLUMBER2 기반 정량 성능평가 체계"),
])

# ================= 1차년도 D : ③ 성능진단 =================
del_placeholder(y1_d); set_title(y1_d, "③ Offline Simulation 및 기본성능 진단")
add_body(y1_d, [
    H("■ 장기 Offline Simulation"),
    I("1979–현재 장기 적분 (4종 LSM, ERA5 forcing)"),
    I("spin-up 수렴 확인(토양수분 smcl, 토양온도 t_soil) → production run"),
    B(),
    H("■ 모델-관측 비교 진단"),
    I("토양수분/온도·식생(LAI/GPP) 기본성능 진단"),
    I("지면-대기 결합강도(Mi, λ)·유출량 기반 해양영향 기초 진단도구 개발"),
    B(),
    H("■ 진행 현황"),
    I("ERA5 1979년 변환 완료, JULES 테스트런 수행 → 전 기간 확장 예정"),
])

# ================= 성과계획 (표) =================
del_placeholder(s_perf)
perf = [
    ["구분", "1차년도(2026) 정량 목표"],
    ["오프라인 테스트베드", "지면-식생모델 4종 (JULES · CLM5 · Noah-MP · LM4+)"],
    ["평가용 자료 DB", "토양수분·토양온도·식생·유출 4부문 (관측/재분석 6종 이상)"],
    ["Offline simulation", "1979–현재 장기적분 (약 40여 년)"],
    ["진단 도구", "지면-대기 결합·해양영향 기초 진단도구 1식"],
    ["학술 성과", "국내외 학회 발표 1건 이상"],
]
t = add_table(s_perf, perf, 1.8, 2.9, 23.0, 8.6, col_w=[6.5, 16.5])
style_table(t, header_sz=22, body_sz=20)
add_body(s_perf, [("※ 정량 목표 수치는 협약서 기준에 맞추어 최종 확정", 17, False, GREY, 0)],
         left=1.8, top=11.9, width=22.0, height=1.0)

# ================= 건의/토의 =================
del_placeholder(s_sugg)
add_body(s_sugg, [
    H("■ 기술 현안"),
    I("JULES MPI 빌드 실패 → 현재 serial 운용, 병렬화 해결 필요"),
    I("ERA5 forcing time coordinate 불일치 → 표준화/재생성 필요"),
    B(),
    H("■ 자료 확보"),
    I("검증용 관측 DB(FLUXNET, GRDC, GLEAM 등) 접근·라이선스"),
    B(),
    H("■ 협업 / 자원"),
    I("해양모델팀과 담수 유출 인터페이스(결합 변수·격자) 규격 협의"),
    I("다중 LSM 장기적분을 위한 계산자원·연구인력 확보"),
])

# ================= reorder =================
desired = [s_title, s_toc, s_plan, plan_b, plan_c,
           s_y1, y1_b, y1_c, y1_d, s_perf, s_sugg, s_thx]
sldIdLst = prs.slides._sldIdLst
id_for = {}
for sldId in list(sldIdLst):
    rId = sldId.get(qn('r:id'))
    id_for[prs.part.rels[rId].target_part] = sldId
for sldId in list(sldIdLst):
    sldIdLst.remove(sldId)
for sl in desired:
    sldIdLst.append(id_for[sl.part])

# page numbers (2-digit, by final position; title has no number box)
for idx, sl in enumerate(desired, start=1):
    if idx == 1:
        continue
    set_pagenum(sl, f"{idx:02d}")

prs.save(OUT)
print("saved:", OUT, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
print("order:", [find(s, "TextBox 6").text_frame.text if find(s,"TextBox 6") else "(title/toc/thx)" for s in desired])
